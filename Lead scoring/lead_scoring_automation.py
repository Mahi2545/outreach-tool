"""
Production-ready B2B Lead Scoring Pipeline (Option B: Python)

What it does
- Reads HubSpot export (CSV/XLSX)
- Uses existing HubSpot fields for Industry, Revenue, Connections when present
- Optionally enriches:
  - Signal 3 (AI hiring) + Signal 6 (Leadership hiring) from company job listings via Proxycurl
  - Signal 5 (CXO in India) via Proxycurl employee search
- Applies the 6-signal scoring formula you specified
- Outputs an enriched Excel with:
  - HOT/WARM/COLD tiers + row coloring
  - Priority rank
  - Golden Lead flag (India CXO + 1st/2nd degree)
  - Signal strength bar (●●●●● / ●●● / ●)
  - AI job description snippets (pitch ammo)

How to run (PowerShell, from repo root)
  & "$env:LocalAppData\\Programs\\Python\\Python312\\python.exe" -m pip install -r requirements.txt
  & "$env:LocalAppData\\Programs\\Python\\Python312\\python.exe" "Lead scoring\\lead_scoring_automation.py" --input "your_hubspot_export.xlsx"

Proxycurl (recommended)
  $env:PROXYCURL_API_KEY="..."
  & "$env:LocalAppData\\Programs\\Python\\Python312\\python.exe" "Lead scoring\\lead_scoring_automation.py" --input "your_hubspot_export.xlsx" --proxycurl-api-key $env:PROXYCURL_API_KEY

Notes
- This script avoids direct LinkedIn crawling. It’s API-first.
- Proxycurl endpoints/fields can vary by plan; adjust `ProxycurlClient` methods if needed.
"""

from __future__ import annotations

import argparse
import os
import json
import hashlib
import re
import time
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from dotenv import load_dotenv
except Exception:  # pragma: no cover
    load_dotenv = None


# ----------------------------
# Config (weights are editable)
# ----------------------------

CONFIG: Dict[str, Any] = {
    "weights": {
        "industry": 15,
        "revenue": 15,
        "ai_hiring": 30,
        "connection": 20,
        "india_cxo": 10,
        "leadership_hiring": 10,
    },
    "tiers": {
        "HOT": (70, 100),
        "WARM": (40, 69),
        "COLD": (0, 39),
    },
    "icp_industries": {
        "financial services",
        "fintech",
        "banking",
        "investment management",
        "insurance",
        "insurance (financial)",
        "payments",
    },
    "ai_keywords": [
        r"\bai\b",
        r"artificial intelligence",
        r"machine learning",
        r"\bml\b",
        r"\bml engineer\b",
        r"data scientist",
        r"\bnlp\b",
        r"\bllm\b",
        r"generative ai",
        r"ai product manager",
        r"ai operations",
        r"automation engineer",
        r"ai analyst",
        r"prompt engineer",
        r"head of ai",
        r"\bvp of ai\b",
        r"chief ai officer",
    ],
    "related_data_keywords": [
        r"data engineer",
        r"analytics",
        r"data analyst",
        r"business intelligence",
    ],
    "leadership_keywords": [
        r"\bvp\b",
        r"\bdirector\b",
        r"\bchief\b",
        r"\bc-suite\b",
        r"\bhead of\b",
        r"managing director",
        r"\bpartner\b",
        r"\bprincipal\b",
        r"\bsvp\b",
        r"\bevp\b",
    ],
    "india_location_keywords": [
        "india",
        "bengaluru",
        "bangalore",
        "mumbai",
        "new delhi",
        "delhi",
        "gurgaon",
        "noida",
        "hyderabad",
        "pune",
        "chennai",
        "kolkata",
        "ahmedabad",
    ],
    "request": {
        "timeout_seconds": 25,
        "sleep_seconds": 0.35,
    },
    "column_candidates": {
        "company_name": ["Company Name", "Company name", "Company"],
        "industry": ["Industry"],
        "revenue": ["Revenue Range", "Annual Revenue", "Revenue"],
        "linkedin_url": [
            "LinkedIn Sales Navigator URL",
            "LinkedIn Sales Nav URL",
            "LinkedIn Sales",
            "LinkedIn URL",
            "LinkedIn Company Page",
            "LinkedIn",
        ],
        "connections": ["Connections", "Connection Degree", "Connection_Degree"],
        "description": ["Company Description", "Description"],
        "employees": ["Employees on LinkedIn", "Employees", "Employee Count"],
        "strategic_priority": ["Strategic Priority", "Priority"],
    },
}


# ----------------------------
# Column inference + parsing
# ----------------------------

def pick_col(df: pd.DataFrame, candidates: List[str]) -> Optional[str]:
    cols = {str(c).strip().lower(): str(c) for c in df.columns}
    for c in candidates:
        k = c.strip().lower()
        if k in cols:
            return cols[k]
    return None


def normalize_text(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


def normalize_industry(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().lower())


def revenue_score_from_value(revenue_value: float) -> int:
    if 50_000_000 <= revenue_value <= 100_000_000:
        return 15
    if 20_000_000 <= revenue_value < 50_000_000:
        return 10
    if 5_000_000 <= revenue_value < 20_000_000:
        return 5
    return 0


def parse_revenue_to_bucket(rev_raw: str) -> Tuple[str, int]:
    text = rev_raw.strip()
    if not text:
        return ("", 0)

    # "$20M - $50M" style
    m = re.search(
        r"\$?\s*([\d.]+)\s*([mk]?)\s*[-–]\s*\$?\s*([\d.]+)\s*([mk]?)",
        text,
        flags=re.I,
    )
    if m:
        lo, lo_s, hi, hi_s = m.group(1), m.group(2), m.group(3), m.group(4)
        lo_val = float(lo) * (1_000_000 if lo_s.lower() == "m" else 1_000 if lo_s.lower() == "k" else 1)
        hi_val = float(hi) * (1_000_000 if hi_s.lower() == "m" else 1_000 if hi_s.lower() == "k" else 1)
        avg = (lo_val + hi_val) / 2
        return (f"${int(lo_val/1_000_000)}M - ${int(hi_val/1_000_000)}M", revenue_score_from_value(avg))

    # numeric revenue
    num = re.sub(r"[^\d.]", "", text)
    if num:
        try:
            val = float(num)
            # heuristic: small values may represent millions
            if val < 1_000:
                val = val * 1_000_000
            return (text, revenue_score_from_value(val))
        except Exception:
            pass

    return (text, 0)


def connection_score(conn: str) -> Tuple[str, int]:
    c = conn.strip().lower()
    if c.startswith("1") or "1st" in c:
        return ("1st", 20)
    if c.startswith("2") or "2nd" in c:
        return ("2nd", 15)
    if c.startswith("3") or "3rd" in c:
        return ("3rd", 0)
    if c in {"none", ""}:
        return ("None", 0)
    return (conn, 0)


def tier_from_score(total: int) -> str:
    if total >= CONFIG["tiers"]["HOT"][0]:
        return "🔴 HOT"
    if total >= CONFIG["tiers"]["WARM"][0]:
        return "🟡 WARM"
    return "🟢 COLD"


def signal_strength_bar(tier: str) -> str:
    if "HOT" in tier:
        return "●●●●●"
    if "WARM" in tier:
        return "●●●"
    return "●"


def contains_any(patterns: List[str], text: str) -> bool:
    for p in patterns:
        if re.search(p, text, flags=re.I):
            return True
    return False


# ----------------------------
# Proxycurl enrichment adapters
# ----------------------------

@dataclass
class ProxycurlClient:
    api_key: str
    timeout_seconds: int = 25

    @property
    def headers(self) -> Dict[str, str]:
        return {"Authorization": f"Bearer {self.api_key}"}

    def get_company_jobs(self, linkedin_company_url: str) -> List[Dict[str, Any]]:
        # Adjust endpoint to your plan if needed
        endpoint = "https://nubela.co/proxycurl/api/linkedin/company/job"
        resp = requests.get(
            endpoint,
            headers=self.headers,
            params={"url": linkedin_company_url},
            timeout=self.timeout_seconds,
        )
        resp.raise_for_status()
        data = resp.json()
        return data.get("jobs") or data.get("job_listings") or []

    def search_cxo(self, linkedin_company_url: str) -> Optional[Dict[str, Any]]:
        # Adjust endpoint to your plan if needed
        endpoint = "https://nubela.co/proxycurl/api/linkedin/company/employees"
        resp = requests.get(
            endpoint,
            headers=self.headers,
            params={
                "url": linkedin_company_url,
                "role_search": "CEO OR CTO OR CPO OR COO OR CISO OR \"Chief Data\" OR \"Head of\" OR VP",
                "page_size": 10,
            },
            timeout=self.timeout_seconds,
        )
        resp.raise_for_status()
        data = resp.json()
        people = data.get("employees") or data.get("results") or []
        if not people:
            return None
        p = people[0]
        return {
            "name": p.get("full_name") or p.get("name") or "",
            "title": p.get("title") or p.get("occupation") or "",
            "location": p.get("location") or "",
        }


# ----------------------------
# Coresignal enrichment adapters (trial-friendly)
# ----------------------------

@dataclass
class CoresignalClient:
    api_key: str
    timeout_seconds: int = 25
    base_url: str = "https://api.coresignal.com/cdapi"
    cache_path: str = ".coresignal_cache.json"

    def _headers(self) -> Dict[str, str]:
        return {"apikey": self.api_key, "accept": "application/json", "Content-Type": "application/json"}

    def _cache_key(self, method: str, url: str, payload: Optional[Dict[str, Any]]) -> str:
        blob = json.dumps({"m": method, "u": url, "p": payload or {}}, sort_keys=True).encode("utf-8")
        return hashlib.sha256(blob).hexdigest()

    def _load_cache(self) -> Dict[str, Any]:
        try:
            with open(self.cache_path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}

    def _save_cache(self, cache: Dict[str, Any]) -> None:
        with open(self.cache_path, "w", encoding="utf-8") as f:
            json.dump(cache, f)

    def _post_cached(self, path: str, payload: Dict[str, Any]) -> Dict[str, Any]:
        url = f"{self.base_url}{path}"
        cache = self._load_cache()
        key = self._cache_key("POST", url, payload)
        if key in cache:
            return cache[key]

        resp = requests.post(url, headers=self._headers(), json=payload, timeout=self.timeout_seconds)
        resp.raise_for_status()
        data = resp.json()
        cache[key] = data
        self._save_cache(cache)
        return data

    def _get_cached(self, path: str) -> Dict[str, Any]:
        url = f"{self.base_url}{path}"
        cache = self._load_cache()
        key = self._cache_key("GET", url, None)
        if key in cache:
            return cache[key]

        resp = requests.get(url, headers=self._headers(), timeout=self.timeout_seconds)
        resp.raise_for_status()
        data = resp.json()
        cache[key] = data
        self._save_cache(cache)
        return data

    def get_company_jobs(self, linkedin_company_url: str, limit: int = 25) -> List[Dict[str, Any]]:
        """
        Uses Base Jobs API:
        - POST /v2/job_base/search/filter/preview to retrieve a small set of job IDs (cheaper)
        - GET  /v2/job_base/collect/{job_id} for top N IDs to get titles + description snippets
        """
        preview = self._post_cached(
            "/v2/job_base/search/filter/preview",
            {
                "company_professional_network_url": linkedin_company_url,
                "application_active": True,
                "deleted": False,
            },
        )
        ids = preview.get("ids") or preview.get("data") or []
        ids = ids[: max(0, limit)]
        jobs: List[Dict[str, Any]] = []
        for job_id in ids:
            try:
                job = self._get_cached(f"/v2/job_base/collect/{job_id}")
                jobs.append(job)
            except Exception:
                continue
        return jobs

    def search_cxo(self, linkedin_company_url: str) -> Optional[Dict[str, Any]]:
        """
        Uses Base Employee API:
        - POST /v2/employee_base/search/filter/preview to retrieve IDs for active experiences at a company
        - GET  /v2/employee_base/collect/{employee_id} to fetch profile incl. location/headline
        """
        preview = self._post_cached(
            "/v2/employee_base/search/filter/preview",
            {
                "experience_company_professional_network_url": linkedin_company_url,
                "active_experience": True,
                "experience_title": "(CEO) OR (CTO) OR (CPO) OR (COO) OR (CISO) OR (Chief Data) OR (Head of) OR (VP)",
                "deleted": False,
            },
        )
        ids = preview.get("ids") or preview.get("data") or []
        if not ids:
            return None

        # Collect just the top 1 (trial-friendly)
        try:
            p = self._get_cached(f"/v2/employee_base/collect/{ids[0]}")
        except Exception:
            return None

        return {
            "name": p.get("full_name") or "",
            "title": p.get("headline") or p.get("experience", [{}])[0].get("title", ""),
            "location": p.get("location") or "",
        }


# ----------------------------
# Signal scoring
# ----------------------------

def industry_filter(industry_raw: str) -> Tuple[str, int, bool]:
    industry_norm = normalize_industry(industry_raw)
    matches = any(icp in industry_norm for icp in CONFIG["icp_industries"])
    if not matches:
        return (industry_raw, 0, True)  # disqualify
    return (industry_raw, CONFIG["weights"]["industry"], False)


def score_ai_and_leadership_from_jobs(jobs: List[Dict[str, Any]]) -> Dict[str, Any]:
    titles: List[str] = []
    snippets: List[str] = []
    for j in jobs:
        title = normalize_text(j.get("job_title") or j.get("title") or "")
        desc = normalize_text(
            j.get("job_description")
            or j.get("description")
            or j.get("job_description_snippet")
            or ""
        )
        if title:
            titles.append(title)
        if desc:
            snippets.append(desc[:240])

    ai_titles = [t for t in titles if contains_any(CONFIG["ai_keywords"], t)]
    related_titles = [t for t in titles if contains_any(CONFIG["related_data_keywords"], t)]
    leadership_titles = [t for t in titles if contains_any(CONFIG["leadership_keywords"], t)]

    # Signal 3 scoring
    if len(ai_titles) >= 2:
        ai_score = 30
    elif len(ai_titles) == 1:
        ai_score = 20
    elif len(related_titles) >= 1:
        ai_score = 10
    else:
        ai_score = 0

    # Signal 6 scoring
    if len(leadership_titles) >= 2:
        leadership_score = 10
    elif len(leadership_titles) == 1:
        leadership_score = 5
    else:
        leadership_score = 0

    return {
        "AI_Hiring_Roles": "; ".join(sorted(set(ai_titles)))[:1000],
        "AI_Hiring_Score": ai_score,
        "AI_Job_Snippets": " | ".join(snippets[:5])[:2000],
        "Leadership_Hiring": "; ".join(sorted(set(leadership_titles)))[:1000],
        "Leadership_Score": leadership_score,
    }


def india_cxo_score(cxo_location: str) -> Tuple[str, int, str]:
    loc = cxo_location.strip()
    loc_l = loc.lower()
    is_india = any(k in loc_l for k in CONFIG["india_location_keywords"])
    return ("Yes" if is_india else "No", (CONFIG["weights"]["india_cxo"] if is_india else 0), loc)


def build_signal_summary(parts: List[str]) -> str:
    return ", ".join([p for p in parts if p])[:220]


# ----------------------------
# Excel formatting
# ----------------------------

def apply_excel_formatting(path: str) -> None:
    wb = load_workbook(path)
    ws = wb.active

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fills = {
        "🔴 HOT": PatternFill("solid", fgColor="FEE2E2"),
        "🟡 WARM": PatternFill("solid", fgColor="FEF3C7"),
        "🟢 COLD": PatternFill("solid", fgColor="DCFCE7"),
    }

    headers = [c.value for c in ws[1]]
    if "Tier" in headers:
        tier_idx = headers.index("Tier") + 1
        for r in range(2, ws.max_row + 1):
            tier_val = ws.cell(row=r, column=tier_idx).value or ""
            fill = fills.get(str(tier_val), None)
            if fill:
                for c in range(1, ws.max_column + 1):
                    ws.cell(row=r, column=c).fill = fill

    wb.save(path)


# ----------------------------
# IO + main pipeline
# ----------------------------

def read_input(path: str) -> pd.DataFrame:
    if path.lower().endswith(".csv"):
        return pd.read_csv(path)
    return pd.read_excel(path)


def main() -> None:
    # Load optional .env next to this script (Lead scoring/.env)
    if load_dotenv is not None:
        env_path = os.path.join(os.path.dirname(__file__), ".env")
        if os.path.exists(env_path):
            load_dotenv(env_path, override=False)

    parser = argparse.ArgumentParser(description="B2B Lead Scoring Pipeline")
    parser.add_argument("--input", required=True, help="HubSpot export CSV/XLSX")
    parser.add_argument("--output", default="", help="Output XLSX path")
    parser.add_argument("--proxycurl-api-key", default=os.getenv("PROXYCURL_API_KEY", ""), help="Proxycurl API key (or set env PROXYCURL_API_KEY)")
    parser.add_argument("--coresignal-api-key", default=os.getenv("CORESIGNAL_API_KEY", ""), help="Coresignal API key (or set env CORESIGNAL_API_KEY)")
    parser.add_argument("--max-companies", type=int, default=0, help="Limit rows for testing")
    parser.add_argument("--tail", type=int, default=0, help="Process only the last N rows (useful for verification)")
    parser.add_argument("--sleep", type=float, default=CONFIG["request"]["sleep_seconds"], help="Sleep between API calls")
    parser.add_argument("--coresignal-cache", default=".coresignal_cache.json", help="Cache file path to reduce trial credit usage")
    args = parser.parse_args()

    df = read_input(args.input)
    if args.tail and args.tail > 0:
        df = df.tail(args.tail).copy()
    if args.max_companies and args.max_companies > 0:
        df = df.head(args.max_companies).copy()

    col_company = pick_col(df, CONFIG["column_candidates"]["company_name"])
    col_industry = pick_col(df, CONFIG["column_candidates"]["industry"])
    col_revenue = pick_col(df, CONFIG["column_candidates"]["revenue"])
    col_linkedin = pick_col(df, CONFIG["column_candidates"]["linkedin_url"])
    col_conn = pick_col(df, CONFIG["column_candidates"]["connections"])

    if not col_company:
        raise ValueError("Could not find a Company Name column. Add/rename it to 'Company Name'.")

    proxycurl: Optional[ProxycurlClient] = None
    if args.proxycurl_api_key:
        proxycurl = ProxycurlClient(api_key=args.proxycurl_api_key, timeout_seconds=CONFIG["request"]["timeout_seconds"])

    coresignal: Optional[CoresignalClient] = None
    if args.coresignal_api_key:
        coresignal = CoresignalClient(
            api_key=args.coresignal_api_key,
            timeout_seconds=CONFIG["request"]["timeout_seconds"],
            cache_path=args.coresignal_cache,
        )

    out_rows: List[Dict[str, Any]] = []
    start = time.time()

    for idx, row in df.iterrows():
        company_name = normalize_text(row.get(col_company))
        industry_raw = normalize_text(row.get(col_industry)) if col_industry else ""
        revenue_raw = normalize_text(row.get(col_revenue)) if col_revenue else ""
        linkedin_url = normalize_text(row.get(col_linkedin)) if col_linkedin else ""
        conn_raw = normalize_text(row.get(col_conn)) if col_conn else ""

        # Signal 1
        industry_verified, industry_score, disqualify = industry_filter(industry_raw)
        if disqualify:
            continue

        # Signal 2
        revenue_range_norm, revenue_score = parse_revenue_to_bucket(revenue_raw)

        # Signal 4
        conn_degree, conn_score_val = connection_score(conn_raw)

        # Signals 3/6 + 5 (optional enrichment)
        jobs_payload: List[Dict[str, Any]] = []
        cxo: Dict[str, Any] = {"name": "", "title": "", "location": ""}

        # Enrichment preference: Coresignal first (trial-friendly via preview + cache), else Proxycurl
        if coresignal and linkedin_url:
            try:
                jobs_payload = coresignal.get_company_jobs(linkedin_url, limit=25)
            except Exception:
                jobs_payload = []
            time.sleep(args.sleep)
            try:
                cxo_found = coresignal.search_cxo(linkedin_url)
                if cxo_found:
                    cxo = cxo_found
            except Exception:
                pass
            time.sleep(args.sleep)
        elif proxycurl and linkedin_url:
            try:
                jobs_payload = proxycurl.get_company_jobs(linkedin_url)
            except Exception:
                jobs_payload = []
            time.sleep(args.sleep)
            try:
                cxo_found = proxycurl.search_cxo(linkedin_url)
                if cxo_found:
                    cxo = cxo_found
            except Exception:
                pass
            time.sleep(args.sleep)

        jobs_scored = score_ai_and_leadership_from_jobs(jobs_payload) if jobs_payload else {
            "AI_Hiring_Roles": "",
            "AI_Hiring_Score": 0,
            "AI_Job_Snippets": "",
            "Leadership_Hiring": "",
            "Leadership_Score": 0,
        }

        india_cxo, india_score, cxo_location = india_cxo_score(normalize_text(cxo.get("location", "")))

        total = int(
            industry_score
            + revenue_score
            + jobs_scored["AI_Hiring_Score"]
            + conn_score_val
            + india_score
            + jobs_scored["Leadership_Score"]
        )
        tier = tier_from_score(total)

        # Additional requirement: Golden Lead = India CXO + 1st/2nd degree
        golden = "Yes" if (india_cxo == "Yes" and conn_degree in {"1st", "2nd"}) else "No"

        summary_parts: List[str] = []
        if jobs_scored["AI_Hiring_Score"] > 0:
            summary_parts.append(f"AI roles: {jobs_scored['AI_Hiring_Roles'][:80]}")
        if conn_degree in {"1st", "2nd"}:
            summary_parts.append(f"{conn_degree} degree")
        if india_cxo == "Yes":
            summary_parts.append("India CXO")
        if revenue_score > 0:
            summary_parts.append(revenue_range_norm or revenue_raw)

        out_rows.append(
            {
                "Company_Name": company_name,
                "Industry": industry_verified,
                "Revenue_Range": revenue_range_norm or revenue_raw,
                "LinkedIn_URL": linkedin_url,
                "AI_Hiring_Roles": jobs_scored["AI_Hiring_Roles"],
                "AI_Job_Snippets": jobs_scored["AI_Job_Snippets"],
                "AI_Hiring_Score": jobs_scored["AI_Hiring_Score"],
                "Connection_Degree": conn_degree,
                "Connection_Score": conn_score_val,
                "CXO_Name": normalize_text(cxo.get("name", "")),
                "CXO_Title": normalize_text(cxo.get("title", "")),
                "CXO_Location": cxo_location,
                "India_CXO": india_cxo,
                "India_CXO_Score": india_score,
                "Leadership_Hiring": jobs_scored["Leadership_Hiring"],
                "Leadership_Score": jobs_scored["Leadership_Score"],
                "Revenue_Score": revenue_score,
                "Industry_Score": industry_score,
                "Total_Score": total,
                "Tier": tier,
                "Golden_Lead": golden,
                "Priority_Rank": 0,  # filled after sorting
                "Signal_Strength_Bar": signal_strength_bar(tier),
                "Signal_Summary": build_signal_summary(summary_parts),
            }
        )

        if (idx + 1) % 25 == 0:
            elapsed = time.time() - start
            print(f"Processed input rows: {idx + 1} (elapsed {elapsed:.1f}s)", flush=True)

    out = pd.DataFrame(out_rows)
    if out.empty:
        raise RuntimeError("No companies passed the Industry filter. Check Industry values/ICP list.")

    out["__golden_sort"] = (out["Golden_Lead"] == "Yes").astype(int)
    out = (
        out.sort_values(["__golden_sort", "Total_Score"], ascending=[False, False])
        .drop(columns=["__golden_sort"])
        .reset_index(drop=True)
    )
    out["Priority_Rank"] = out.index + 1

    output_path = args.output.strip() or os.path.splitext(args.input)[0] + ".lead_scored.xlsx"
    out.to_excel(output_path, index=False)
    apply_excel_formatting(output_path)

    print(f"Saved: {output_path}")
    hot = out[out["Tier"].str.contains("HOT", na=False)].head(20)
    print(f"HOT leads: {len(hot)} (top {min(20, len(hot))})")
    for _, r in hot.iterrows():
        print(f"- #{int(r['Priority_Rank'])} {r['Company_Name']} ({int(r['Total_Score'])}) — {r['Signal_Summary']}")


if __name__ == "__main__":
    main()