import pandas as pd
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

FINSERV_INDUSTRIES = {"Financial Services","Banking","Insurance","Investment Banking","Capital Markets"}

def score_industry(val):
    if pd.isna(val): return 0
    v = str(val).strip()
    if v in FINSERV_INDUSTRIES: return 25
    if v == "Investment Management": return 10
    return 0

def score_revenue(val):
    if pd.isna(val): return 0
    ICP = {"$5M - $10M","$10M - $20M","$20M - $50M","$50M - $100M"}
    return 20 if str(val).strip() in ICP else 0

def score_ai_hiring(val):
    if pd.isna(val): return 0
    return 20 if str(val).strip().lower() == "yes" else 0

def score_network(val):
    if pd.isna(val): return 0
    v = str(val)
    if v == "-": return 0
    if "1st" in v: return 15
    if "2nd" in v: return 10
    return 0

def score_india(val):
    if pd.isna(val): return 0
    v = str(val).lower()
    if v == "-": return 0
    return 12 if "india" in v else 0

def score_leadership(val):
    if pd.isna(val): return 0
    v = str(val).lower()
    if "senior leadership hire" not in v: return 0
    m = re.match(r"(\d+)", v.strip())
    n = int(m.group(1)) if m else 1
    if n >= 3: return 8
    if n == 2: return 6
    return 4

def calculate_tier(score):
    if score >= 75: return "Tier 1 — Hot"
    if score >= 50: return "Tier 2 — Warm"
    if score >= 25: return "Tier 3 — Cold"
    return "Tier 4 — Disqualified"

def recommended_action(tier):
    return {
        "Tier 1 — Hot": "Priority outreach this week — warm intro or direct contact within 48hrs",
        "Tier 2 — Warm": "2-touch nurture sequence — contact within 2 weeks",
        "Tier 3 — Cold": "Enrich further, add to monthly content drip",
        "Tier 4 — Disqualified": "Do not contact — revisit if signals change"
    }[tier]

def generate_talk_track(row):
    parts = []
    name = row["Company Name"]
    industry = row["Industry"]
    rev = row["Revenue Range"]

    if row["s3_ai_hiring"] > 0 and row["s6_leadership_hire"] > 0:
        parts.append(f"We noticed {name} is actively expanding its AI capabilities and building out senior leadership — exactly the growth stage where our platform delivers outsized ROI.")
    elif row["s3_ai_hiring"] > 0:
        parts.append(f"We saw {name} is hiring for AI-related roles — we work with {industry} firms at this exact inflection point to accelerate time-to-value.")
    elif row["s6_leadership_hire"] > 0:
        parts.append(f"{name} is investing in senior leadership — new leaders typically reset vendor priorities in the first 90 days, and we'd love to be on your radar early.")

    if row["s5_india_cxo"] > 0:
        parts.append("Given your India-based leadership, we can offer local context and timezone-aligned support.")

    if not parts:
        parts.append(f"We work with {industry} companies in the {rev} range on AI-driven operations — would love to share what's worked for similar firms.")

    if row["s4_network"] >= 10:
        parts.append("Our mutual connection would be happy to make a warm introduction.")

    sp = str(row.get("Strategic Priority",""))
    if sp and sp not in ("-","nan",""):
        parts.append(f"Your focus on '{sp}' aligns closely with what we help teams achieve.")

    return " ".join(parts)

def load_and_score(input_file):
    df = pd.read_excel(input_file)
    df["s1_industry"]       = df["Industry"].apply(score_industry)
    df["s2_revenue"]        = df["Revenue Range"].apply(score_revenue)
    df["s3_ai_hiring"]      = df["Hiring on LinkedIn"].apply(score_ai_hiring)
    df["s4_network"]        = df["Signal 5 (Connection)"].apply(score_network)
    df["s5_india_cxo"]      = df["Signal 6(Location)"].apply(score_india)
    df["s6_leadership_hire"]= df["Senior Leadership Hires"].apply(score_leadership)
    df["Total Score"] = (df["s1_industry"]+df["s2_revenue"]+df["s3_ai_hiring"]+
                         df["s4_network"]+df["s5_india_cxo"]+df["s6_leadership_hire"])
    df["Tier"] = df["Total Score"].apply(calculate_tier)
    df["Recommended Action"] = df["Tier"].apply(recommended_action)
    df["Talk Track"] = df.apply(generate_talk_track, axis=1)
    df = df.sort_values("Total Score", ascending=False).reset_index(drop=True)
    return df

TIER_FILLS = {
    "Tier 1 — Hot":        PatternFill("solid", fgColor="C6EFCE"),
    "Tier 2 — Warm":       PatternFill("solid", fgColor="FFEB9C"),
    "Tier 3 — Cold":       PatternFill("solid", fgColor="FFCC99"),
    "Tier 4 — Disqualified": PatternFill("solid", fgColor="FFFFFF"),
}

def autofit(ws, max_width=60):
    for col in ws.columns:
        w = 10
        for cell in col:
            try:
                if cell.value:
                    w = max(w, min(len(str(cell.value))+2, max_width))
            except: pass
        ws.column_dimensions[get_column_letter(col[0].column)].width = w

def write_scored_leads(df, out_path):
    df.to_excel(out_path, index=False, sheet_name="All Companies")
    wb = load_workbook(out_path)
    ws = wb.active
    hdr_fill = PatternFill("solid", fgColor="404040")
    hdr_font = Font(bold=True, color="FFFFFF", name="Arial")
    for cell in ws[1]:
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
    tier_col = [c.column for c in ws[1] if c.value == "Tier"]
    tier_col = tier_col[0] if tier_col else None
    for row in ws.iter_rows(min_row=2):
        tier_val = None
        if tier_col:
            tier_val = row[tier_col-1].value
        fill = TIER_FILLS.get(tier_val, PatternFill("solid", fgColor="FFFFFF"))
        for cell in row:
            cell.fill = fill
            cell.font = Font(name="Arial", size=10)
    ws.freeze_panes = "A2"
    autofit(ws)
    wb.save(out_path)

def write_boss_report(df, out_path):
    boss_cols = [
        "Company Name","Industry","Revenue Range","Total Score","Tier",
        "Hiring on LinkedIn","Signal 5 (Connection)","Signal 6(Location)",
        "Senior Leadership Hires","Strategic Priority","Fit",
        "Recommended Action","Talk Track","LinkedIn Sales Nav URL"
    ]
    rename_map = {
        "Signal 5 (Connection)": "Network Degree",
        "Signal 6(Location)":    "CXO Location"
    }
    top = df[df["Tier"].isin(["Tier 1 — Hot","Tier 2 — Warm"])][boss_cols].copy()
    top = top.rename(columns=rename_map)

    summary_data = {
        "Metric": [
            "Total companies analyzed",
            "Tier 1 — Hot (score ≥75)",
            "Tier 2 — Warm (score 50–74)",
            "Tier 3 — Cold (score 25–49)",
            "Tier 4 — Disqualified (<25)",
            "Companies with AI hiring signal",
            "Companies with India CXO",
            "Companies in primary FS industries",
            "Companies with 1st degree connection",
        ],
        "Count": [
            len(df),
            len(df[df["Tier"]=="Tier 1 — Hot"]),
            len(df[df["Tier"]=="Tier 2 — Warm"]),
            len(df[df["Tier"]=="Tier 3 — Cold"]),
            len(df[df["Tier"]=="Tier 4 — Disqualified"]),
            len(df[df["s3_ai_hiring"]>0]),
            len(df[df["s5_india_cxo"]>0]),
            len(df[df["s1_industry"]==25]),
            len(df[df["s4_network"]==15]),
        ]
    }
    summary_df = pd.DataFrame(summary_data)
    summary_df["% of Total"] = (summary_df["Count"]/len(df)*100).round(1).astype(str)+"%"

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        top.to_excel(writer, index=False, sheet_name="Outreach Priority")
        summary_df.to_excel(writer, index=False, sheet_name="Score Summary")

    wb = load_workbook(out_path)

    # Format Outreach Priority sheet
    ws = wb["Outreach Priority"]
    navy_fill = PatternFill("solid", fgColor="1F3864")
    hdr_font  = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    thick_bot = Border(bottom=Side(style="thick", color="000000"))
    for cell in ws[1]:
        cell.fill = navy_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thick_bot

    tier_col_idx = None
    name_col_idx = None
    score_col_idx = None
    talk_col_idx = None
    for c in ws[1]:
        if c.value == "Tier": tier_col_idx = c.column
        if c.value == "Company Name": name_col_idx = c.column
        if c.value == "Total Score": score_col_idx = c.column
        if c.value == "Talk Track": talk_col_idx = c.column

    for row in ws.iter_rows(min_row=2):
        tier_val = row[tier_col_idx-1].value if tier_col_idx else None
        row_fill = TIER_FILLS.get(tier_val, PatternFill("solid", fgColor="FFFFFF"))
        for cell in row:
            cell.fill = row_fill
            cell.font = Font(name="Arial", size=10,
                             bold=(cell.column == name_col_idx))
            if cell.column == score_col_idx:
                cell.alignment = Alignment(horizontal="center")
                cell.font = Font(name="Arial", size=10, bold=True)
            if cell.column == talk_col_idx:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    if talk_col_idx:
        ws.column_dimensions[get_column_letter(talk_col_idx)].width = 55
    if score_col_idx:
        ws.column_dimensions[get_column_letter(score_col_idx)].width = 12
    if name_col_idx:
        ws.column_dimensions[get_column_letter(name_col_idx)].width = 35

    for i, col in enumerate(ws.columns, 1):
        if i not in [talk_col_idx, name_col_idx, score_col_idx]:
            w = 10
            for cell in col:
                try:
                    if cell.value: w = max(w, min(len(str(cell.value))+2, 45))
                except: pass
            ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    # Format Score Summary sheet
    ws2 = wb["Score Summary"]
    for cell in ws2[1]:
        cell.fill = navy_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
    autofit(ws2, max_width=50)

    wb.save(out_path)

if __name__ == "__main__":
    input_file = "/mnt/user-data/uploads/NY_lead_sheet_.xlsx"
    df = load_and_score(input_file)

    # Sanity check
    garp = df[df["Company Name"].str.contains("GARP", na=False)].iloc[0]
    solytics = df[df["Company Name"].str.contains("Solytics", na=False)].iloc[0]
    print("=== SANITY CHECK ===")
    print(f"GARP total: {garp['Total Score']} (expected 71) | Tier: {garp['Tier']}")
    print(f"Solytics total: {solytics['Total Score']} (expected 91) | Tier: {solytics['Tier']}")
    print()
    print("=== TIER DISTRIBUTION ===")
    print(df["Tier"].value_counts().to_string())

    write_scored_leads(df, "/home/claude/scored_leads.xlsx")
    write_boss_report(df, "/home/claude/boss_report.xlsx")
    print("\nDone. Files written.")
